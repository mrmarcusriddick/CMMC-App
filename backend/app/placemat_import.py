"""Read Microsoft Product Placemat workbooks without executing VBA macros."""
import argparse
import hashlib
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from .database import Base, SessionLocal, engine
from .models import MicrosoftPlacematRelease, MicrosoftPracticeMapping


class PlacematValidationError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def services(value: str | None) -> list[str]:
    if not value:
        return []
    return [item.strip().replace("\xa0", " ") for item in value.splitlines() if item.strip()]


def extract_mappings(path: Path) -> list[dict]:
    if path.suffix.lower() != ".xlsm":
        raise PlacematValidationError("Microsoft Product Placemat must be supplied as its original .xlsm workbook.")
    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
    required = {"CMMC Levels", "Control Information", "Statements"}
    if not required.issubset(workbook.sheetnames):
        raise PlacematValidationError("Workbook does not contain the required Product Placemat data sheets.")
    level_two = {row[0] for row in workbook["CMMC Levels"].iter_rows(min_row=2, values_only=True) if row[3] == "X" and row[0]}
    if len(level_two) != 110:
        raise PlacematValidationError(f"Expected 110 Level 2 practices; found {len(level_two)}.")
    statements = {row[1]: row[2] for row in workbook["Statements"].iter_rows(min_row=2, values_only=True) if row[1] and row[2]}
    rows = []
    for row in workbook["Control Information"].iter_rows(min_row=3, values_only=True):
        legacy, practice, statement = row[1], row[2], row[3]
        if practice not in level_two:
            continue
        rows.append({"legacy_identifier": legacy, "practice_identifier": practice, "requirement_statement": statement or "", "primary_services": services(row[4]), "secondary_services": services(row[5]), "commercial_responsibility": row[6], "gcc_high_responsibility": row[7], "nist_800_171_mapping": row[8], "nist_800_53_mapping": row[10] or row[9], "implementation_statement": statements.get(practice)})
    if len(rows) != 110 or len({row["practice_identifier"] for row in rows}) != 110:
        raise PlacematValidationError("Control Information does not provide one unique mapping for each Level 2 practice.")
    return rows


def import_placemat(session: Session, path: Path) -> MicrosoftPlacematRelease:
    source_hash = sha256_file(path)
    existing = session.scalar(select(MicrosoftPlacematRelease).where(MicrosoftPlacematRelease.source_sha256 == source_hash))
    if existing:
        return existing
    mappings = extract_mappings(path)
    release = MicrosoftPlacematRelease(title="Microsoft Product Placemat for CMMC", version_label="Preview September 2024", source_filename=path.name, source_sha256=source_hash, status="IMPORTED_UNREVIEWED")
    session.add(release)
    session.flush()
    for item in mappings:
        session.add(MicrosoftPracticeMapping(release_id=release.id, **item))
    session.commit()
    return release


def main() -> None:
    parser = argparse.ArgumentParser(description="Import a Microsoft Product Placemat as informational coverage guidance.")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    Base.metadata.create_all(engine)
    with SessionLocal() as session:
        release = import_placemat(session, args.workbook)
        result = (release.title, release.version_label, release.id, release.status)
    print(f"Imported {result[0]} ({result[1]}): {result[2]} [{result[3]}]")


if __name__ == "__main__":
    main()
